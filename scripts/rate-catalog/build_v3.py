"""Build v3 charge-builder model: JSON + Excel + HTML mockup."""
import os, sys, json, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import model_v3 as M

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
OUTDIR = f"{ROOT}/workingfolder-output"
import source as SRC  # BillingItem_Usage1001 -Final.xlsx, read by column name
usage, master = SRC.usage, SRC.master

charges = {c["id"]: (cat, c) for cat in M.CATEGORIES for c in cat["charges"]}

# ------------------------------------------------------------------ code -> charge
code_map = {}
for cid, (_, c) in charges.items():
    for code in c.get("codes", []) + M.EXTRA_CODES.get(cid, []):
        assert cid in charges
        code_map.setdefault(code, cid)
for code in M.INTERNAL:
    code_map.setdefault(code, "INTERNAL")
for code, m in master.items():
    if code in code_map:
        continue
    trig = usage[code]["trigger"] if code in usage else m["trigger"]
    if trig in M.TRIGGER_TO_CHARGE:
        code_map[code] = M.TRIGGER_TO_CHARGE[trig]
unmapped = sorted(set(master) - set(code_map))
unknown = sorted(c for c in code_map if c not in master)
assert not unmapped, f"unmapped: {unmapped}"
assert not unknown, f"unknown: {unknown}"
by_charge = collections.defaultdict(list)
for code, cid in code_map.items():
    by_charge[cid].append(code)
def inv_lines(cid):
    return sum(usage[c]["freq"] for c in by_charge[cid] if c in usage)

# ------------------------------------------------------------------ JSON
model = {"version": "v3", "categories": [], "proposal": {"sections": M.PROPOSAL_SECTIONS, "categoryToSection": M.CATEGORY_TO_SECTION,
         "materials": M.MATERIALS, "mergeGroups": M.MERGE_GROUPS},
         "defaultPreset": M.DEFAULT_PRESET}
assert all(k in charges for k in M.DEFAULT_PRESET), [k for k in M.DEFAULT_PRESET if k not in charges]

# ---- mapback tables + system code metadata (names / UOM / condition keys only - no customer data)
import re as _re
def _cond_keys(cond):
    """Condition keys of a usage-sheet Condition cell. Exports separate keys with ';' or ', <Key>:' (older layout)."""
    if not cond:
        return []
    keys = set()
    for seg in _re.split(r"[;\n]", str(cond)):
        for m in _re.finditer(r"(?:^|,)\s*([A-Za-z][A-Za-z ]*?)\s*:", seg):
            keys.add(m.group(1).strip())
    return sorted(k for k in keys if k)
system_codes = {}
for code, m in master.items():
    u = usage.get(code)
    # no description field: billing-system descriptions embed customer names ("Special Customer (...)")
    system_codes[code] = {"name": (u["name"] if u else m["name"]), "uom": (u["uom"] if u else m["uom"]),
                          "category": (u["category"] if u else m["category"]), "keys": _cond_keys(u["cond"]) if u else [],
                          "invoiceLines": u["freq"] if u else 0, "hlCustomers": u["hl_customers"] if u else 0}
def _codes_in(x):
    if isinstance(x, str): return [x]
    if isinstance(x, list): return x
    if isinstance(x, dict): return [c for k, v in x.items() if k != "fixed" for c in _codes_in(v)]  # fixed = condition values
    return []
for cid, units in M.BUILDER_CODES.items():
    c = charges[cid][1]
    for u in c["units"]:
        assert u["id"] in units, f"{cid}.{u['id']} has no code entry"
        for d in u["drivers"]:
            assert d["id"] in units[u["id"]].get("drivers", {}), f"{cid}.{u['id']}.{d['id']} has no code entry"
    for code in _codes_in(units):
        assert code in master, f"unknown code {code}"
for cid, (_, c) in charges.items():
    if c["kind"] == "simple":
        assert (cid in M.SIMPLE_CODES) != (cid in M.NEW_ITEM_NAMES), f"{cid} must be in exactly one of SIMPLE_CODES / NEW_ITEM_NAMES"
    for cd in c.get("conds", []):
        assert cd["id"] in M.COND_SYSTEM, f"factor {cd['id']} has no system condition"
    for u in c.get("units", []):
        for d in u["drivers"]:
            assert d["id"] in M.DRIVER_SYSTEM, f"driver {d['id']} has no system condition"
for code in M.SIMPLE_CODES.values():
    assert code in master, code
for cid, ex in M.SIMPLE_EXTRA.items():
    assert cid in M.SIMPLE_CODES, cid
    for code in ex.get("alt", []):
        assert code in master, f"unknown alternate code {code}"
model["commonItems"] = [{"code": c["code"], "name": c["name"], "uom": c["uom"], "tag": c["tag"], "hlCustomers": c["hl_customers"]} for c in SRC.common]
model["mapback"] = {"builder": M.BUILDER_CODES, "simple": M.SIMPLE_CODES, "simpleExtra": M.SIMPLE_EXTRA, "newItems": M.NEW_ITEM_NAMES, "notes": M.NOTES,
                    "conds": M.COND_SYSTEM, "drivers": M.DRIVER_SYSTEM, "chargeCondOverride": M.CHARGE_COND_OVERRIDE}
model["systemCodes"] = system_codes
import standard_template
model["standardTemplate"] = standard_template.build()
assert all(k in charges for k in model["standardTemplate"]["selections"]), "standard template references unknown charges"

# privacy guard: the catalog is published publicly - fail the build if any price-list customer name / code leaks in
_customers = {x.strip() for r in SRC.price_rows for x in (r["customer"], r["customer_name"]) if isinstance(x, str) and len(x.strip()) >= 6}
_blob = json.dumps(model, ensure_ascii=False)
_leaks = sorted(n for n in _customers if n in _blob) + _re.findall(r"Special\s+Customer[^\"]{0,40}", _blob)
assert not _leaks, f"customer names would be published: {_leaks[:10]}"
for cat in M.CATEGORIES:
    out = {"id": cat["id"], "name": cat["name"], "desc": cat["desc"], "charges": []}
    for c in cat["charges"]:
        cc = {k: val for k, val in c.items() if k not in ("triggers",)}
        cc["codes"] = sorted(by_charge[c["id"]])
        cc["invoiceLines"] = inv_lines(c["id"])
        out["charges"].append(cc)
    model["categories"].append(out)

json.dump(model, open(f"{OUTDIR}/rate-card-model-v3.json", "w"), indent=1)
WEB_DATA = f"{ROOT}/web/src/data"
if os.path.isdir(WEB_DATA):
    json.dump(model, open(f"{WEB_DATA}/catalog.json", "w"), indent=1)


# ------------------------------------------------------------------ HTML mockup
tpl = open(f"{HERE}/mockup_template.html").read()
open(f"{OUTDIR}/Quote Webform Mockup v3.html", "w").write(tpl.replace("/*__MODEL__*/null", json.dumps(model)))

# ------------------------------------------------------------------ Excel
F = "Calibri"
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
ADV = PatternFill("solid", fgColor="F2F2F2")
NEW = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(1, i, c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

def put(ws, r, vals, fill=None, bold=()):
    for i, x in enumerate(vals, 1):
        cell = ws.cell(r, i, x)
        cell.alignment = WRAP
        cell.border = BORDER
        cell.font = Font(name=F, size=9, bold=i in bold)
        if fill:
            cell.fill = fill

def price(x, pct=False):
    if x is None:
        return "enter"
    return f"{x:.0%}" if pct else f"${x:,.2f}"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "README"
ws.column_dimensions["A"].width = 120
readme = [
    ("UNIS Rate Card Model v3 - Charge Builder (webform data model)", 14, True),
    ("Replaces the v2 'one row per price line' layout with the progressive structure of your Offload example.", 10, False),
    ("", 10, False),
    ("STRUCTURE", 11, True),
    ("Category (7): Setup & Systems | Inbound | Outbound | Storage | Returns | Value-Added Services | Other Charges.", 10, False),
    ("Charge: SIMPLE = one price field (unit + benchmark default).  BUILDER = progressive questions (Offload, Put away, Transload, Order processing, Picking, Packing, Loading, Storage, Returns).", 10, False),
    ("Main vs Advanced: main charges are shown; advanced ones are folded under 'Advanced' inside each category.", 10, False),
    ("", 10, False),
    ("BUILDER STEPS (asked in this order, most common first)", 11, True),
    ("1. Rate differs by ... ?  (offload type, ship method, channel B2B/D2C, receipt type, facility, title, carrier ...) - common factors shown, others under 'more factors'.", 10, False),
    ("2. Charge per ...  (container, pallet, case, each, hour ... ) - only units that fit the answers to step 1 are offered.", 10, False),
    ("3. Price varies by ...  one driver per unit: flat, an attribute (container size, pallet size, item size, weight) or a volume tier (case / SKU / pallet count). Volume tiers ask Range vs Incremental.", 10, False),
    ("4. Minimum charge (per receipt / container / load / order / month)   5. Different rate for additional units (first / additional).", 10, False),
    ("Result: a generated rate table with benchmark defaults the user can overwrite. 'Show all options' disables the hiding rules.", 10, False),
    ("", 10, False),
    ("TABS", 11, True),
    ("Charges - every charge with category, main/advanced, type, default price & range (simple) and mapped system codes.", 10, False),
    ("Builder Questions - every question and option of every builder, with show-when rules and benchmark defaults.", 10, False),
    ("Show-Hide Rules - the dependency rules in plain language.", 10, False),
    ("Code Mapping - all 334 system charge codes -> charge.", 10, False),
    ("Benchmark defaults follow the v2 method (UNIS median n>=20 inside industry range > UNIS template > industry midpoint). Industry ranges are model estimates, not a cited survey.", 10, False),
]
for i, (t, s, b) in enumerate(readme, 1):
    c = ws.cell(i, 1, t); c.font = Font(name=F, size=s, bold=b); c.alignment = WRAP

# Charges
ws = wb.create_sheet("Charges")
header(ws, ["Category", "Charge ID", "Charge", "Type", "Main / Advanced", "Channel", "What it covers", "Unit", "Default", "Range", "Basis", "Sep-26 invoice lines", "System codes"],
       [18, 16, 30, 9, 10, 8, 46, 14, 10, 16, 24, 10, 40])
r = 2
for cat in M.CATEGORIES:
    for c in cat["charges"]:
        if c["kind"] == "simple":
            vals = [cat["name"], c["id"], c["name"], "Simple", c["tier"].title(), c["channel"], c["desc"], c["unit"], price(c["default"], c["pct"]),
                    f"{price(c['lo'], c['pct'])} - {price(c['hi'], c['pct'])}", c["basis"]]
        else:
            vals = [cat["name"], c["id"], c["name"], "Builder", c["tier"].title(), c["channel"], c["desc"],
                    " / ".join(u["label"] for u in c["units"]), "see Builder Questions", "", ""]
        put(ws, r, vals + [inv_lines(c["id"]), ", ".join(sorted(by_charge[c["id"]]))],
            fill=NEW if c.get("new") else ADV if c["tier"] == "advanced" else None, bold=(3,))
        r += 1

# Builder questions
ws = wb.create_sheet("Builder Questions")
header(ws, ["Charge", "Step", "Question", "Option", "Shown", "Show only when", "Default price", "Range", "Basis / note"],
       [22, 6, 30, 34, 14, 30, 11, 16, 34])
r = 2
def when_txt(w):
    return "; ".join(f"{k} includes {' or '.join(vs)}" for k, vs in w.items()) if w else ""
for cat in M.CATEGORIES:
    for c in cat["charges"]:
        if c["kind"] != "builder":
            continue
        ws.cell(r, 1, f"{c['id']}  {c['name']}  ({cat['name']})").font = Font(name=F, bold=True, size=10)
        for k in range(1, 10):
            ws.cell(r, k).fill = SEC
        r += 1
        for cd in c["conds"]:
            put(ws, r, [c["name"], "1", "Rate differs by ...?", f"{cd['label']}: {', '.join(cd['values']) or '(free text)'}",
                        "Common" if cd["common"] else "More factors", "", "", "", cd["help"]]); r += 1
        for s in c["settings"]:
            put(ws, r, [c["name"], "1b", s["label"], " | ".join(s["options"]), "Always", "", "", "", "Billing setting - no separate price"]); r += 1
        for u in c["units"]:
            fill = None if u["common"] else ADV
            put(ws, r, [c["name"], "2", "Charge per ...", u["label"], "Common" if u["common"] else "More units", when_txt(u["when"]),
                        price(u["flat"]), f"{price(u['lo'])} - {price(u['hi'])}", u["basis"] + (f". {u['note']}" if u["note"] else "")], fill=fill, bold=(4,)); r += 1
            for d in u["drivers"]:
                vals = ", ".join(f"{x['v']} {price(x['d']) if x['d'] is not None else ''}".strip() for x in d["values"])
                put(ws, r, [c["name"], "3", f"{u['label']}: price varies by ...", f"{d['label']} ({'volume tier' if d['kind'] == 'volume' else 'attribute'}): {vals}",
                            "", when_txt(d["when"]), "", "", ("Asks Range vs Incremental. " if d["calc"] else "") + d["help"]], fill=fill); r += 1
            if u["mins"]:
                put(ws, r, [c["name"], "4", f"{u['label']}: minimum charge?", " / ".join(f"per {m}" + (f" ({price(u['minDefault'][m])})" if m in u["minDefault"] else "") for m in u["mins"]),
                            "", "", "", "", ""], fill=fill); r += 1
            if u["second"]:
                put(ws, r, [c["name"], "5", f"{u['label']}: different rate for additional units?", "Yes / No", "", "", "", "", "First unit / each additional unit"], fill=fill); r += 1

# Show-hide rules
ws = wb.create_sheet("Show-Hide Rules")
header(ws, ["Charge", "If the user ...", "Then hide / skip", "Why"], [22, 48, 54, 50])
rules = [
    ("Any builder", "has not ticked a factor in step 1", "the factor's value list and its columns in the rate table", "Only price dimensions the client needs are asked."),
    ("Any builder", "picks a price driver for a unit (step 3)", "all other drivers of that unit (progressive mode = one driver per unit)", "e.g. Container priced by container size -> no case-count / SKU-count / pallet-size questions."),
    ("Any builder", "picks 'Flat' as the price driver", "all driver questions; one price cell per factor combination", ""),
    ("Any builder", "clicks 'Show all options'", "nothing - every factor, unit and driver is shown and drivers can be combined", "Escape hatch for unusual contracts."),
    ("Any builder / category", "does not open 'more factors' / 'more units' / 'Advanced'", "the less common factors, units and advanced charges stay folded", "Most common content first."),
]
for cat in M.CATEGORIES:
    for c in cat["charges"]:
        if c["kind"] != "builder":
            continue
        for u in c["units"]:
            for k, vs in u["when"].items():
                rules.append((c["name"], f"ticks '{k}' and selects only values outside {', '.join(vs)}", f"unit '{u['label']}'", "Unit does not apply to that freight / order type."))
            for d in u["drivers"]:
                for k, vs in d["when"].items():
                    rules.append((c["name"], f"ticks '{k}' and selects only values outside {', '.join(vs)}", f"driver '{d['label']}' under '{u['label']}'", "Driver does not apply."))
for i, row in enumerate(rules, 2):
    put(ws, i, list(row))

# Code mapping
ws = wb.create_sheet("Code Mapping")
header(ws, ["Charge Code", "System Item Name", "UOM", "Trigger", "Charge ID", "Charge", "Category", "Main / Advanced", "Sep-26 invoice lines", "Status"],
       [26, 40, 12, 22, 16, 30, 18, 10, 10, 22])
def name_of(cid):
    if cid == "INTERNAL":
        return ("Billing adjustment (not quoted)", "Internal", "Internal")
    cat, c = charges[cid]
    return (c["name"], cat["name"], c["tier"].title())
for i, code in enumerate(sorted(master, key=lambda x: (code_map[x] == "INTERNAL", code_map[x], x)), 2):
    cid = code_map[code]
    nm, catn, tier = name_of(cid)
    if code in usage:
        u = usage[code]
        vals = [code, u["name"], u["uom"], u["trigger"], cid, nm, catn, tier, u["freq"], "In use" if u["freq"] > 0 else "Configured - not invoiced"]
    else:
        m = master[code]
        vals = [code, m["name"], m["uom"], m["trigger"], cid, nm, catn, tier, None, "Master list only"]
    put(ws, i, vals, fill=ADV if tier == "Advanced" else None)
ws.auto_filter.ref = f"A1:J{len(master) + 1}"

for w in wb.worksheets:
    w.sheet_view.zoomScale = 110
wb.save(f"{OUTDIR}/UNIS Rate Card Model v3 - Charge Builder.xlsx")

n_simple = sum(1 for _, c in charges.values() if c["kind"] == "simple")
n_main = sum(1 for _, c in charges.values() if c["tier"] == "main")
print(f"categories {len(M.CATEGORIES)} charges {len(charges)} (builder {len(charges) - n_simple}, simple {n_simple}; main {n_main}, advanced {len(charges) - n_main})")
print("codes mapped", len(code_map), "internal", sum(1 for x in code_map.values() if x == "INTERNAL"))
