"""Build the v2 (industry-standard, webform-ready) rate card catalog + review of the Active-Items scope file."""
import os, re, sys, math, json, collections, statistics as st
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import spec_v2 as spec

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
IN = f"{ROOT}/workingfolder-input"
OUT = f"{ROOT}/workingfolder-output/UNIS Rate Card Catalog v2 - Industry Standard & Webform Model.xlsx"
JSON_OUT = f"{ROOT}/workingfolder-output/rate-card-catalog-v2.json"

src = openpyxl.load_workbook(f"{IN}/Billing Items Usage Sep 2026 Final(1).xlsx", read_only=True, data_only=True)
usage = {r[1]: r for r in src["All Billing Items Usage"].iter_rows(min_row=2, values_only=True) if r[1]}
master = {r[1]: r for r in src["All Billing Items"].iter_rows(min_row=2, values_only=True) if r[1]}
pl_by = collections.defaultdict(list)
for r in src["All the customers price list"].iter_rows(min_row=2, values_only=True):
    pl_by[r[5]].append(r)

def num(x):
    try:
        return float(str(x).replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None

def nice(x):
    if x is None:
        return None
    step = 0.05 if x < 1 else 0.25 if x < 10 else 1 if x < 100 else 5 if x < 1000 else 50
    return round(round(x / step) * step, 2)

def report(ln):
    """UNIS evidence from customer price lists for this line."""
    if ln["pct"] and ln["ratetype"] != "Mark Up":
        return None, None, None, 0
    vals = []
    for code in ln["codes"]:
        for r in pl_by.get(code, []):
            if (r[10] or "") != ln["ratetype"]:
                continue
            d = str(r[7] or "").strip()
            if ln["inc"] and not re.search(ln["inc"], d, re.I):
                continue
            if ln["exc"] and re.search(ln["exc"], d, re.I):
                continue
            v = num(r[11])
            if v is not None and v > 0:
                vals.append(v)
    if not vals:
        return None, None, None, 0
    q = st.quantiles(vals, n=4) if len(vals) >= 4 else [min(vals), 0, max(vals)]
    return round(st.median(vals), 2), round(q[0], 2), round(q[2], 2), len(vals)

def benchmark(ln, med, p25, p75, n):
    """Default price + range + basis. Hierarchy: explicit override > UNIS median (n>=20, inside industry range)
    > current UNIS template (inside range) > UNIS median small sample (n>=5, inside range) > industry midpoint."""
    lo, hi = ln["lo"], ln["hi"]
    flags = []
    if lo is None:  # specialty line with no industry benchmark -> UNIS data only
        return med, p25, p75, f"UNIS price-list median (n={n}); range = UNIS P25-P75", flags
    inside = med is not None and lo <= med <= hi
    if med is not None and n >= 20 and not inside:
        flags.append(f"UNIS median {med:g} outside industry range")
    tmpl = ln["d2c"] if ln["ch"] == "D2C" and ln["d2c"] is not None else ln["b2b"] if ln["b2b"] is not None else ln["d2c"]
    if ln["default"] is not None:
        return ln["default"], lo, hi, ln["basis"], flags
    if inside and n >= 20:
        return med, lo, hi, f"UNIS price-list median (n={n})", flags
    if tmpl is not None and lo <= tmpl <= hi:
        return tmpl, lo, hi, "UNIS current template rate", flags
    if inside and n >= 5:
        return med, lo, hi, f"UNIS price-list median (small sample, n={n})", flags
    if tmpl is not None:
        flags.append(f"Template rate {tmpl:g} outside industry range")
    return nice((lo + hi) / 2), lo, hi, "Industry benchmark midpoint", flags

# ------------------------------------------------------------------ assemble lines
rows = []
code_primary = {}
for title, services in spec.SECTIONS:
    for s in services:
        for i, ln in enumerate(s["lines"], 1):
            lid = f"{s['id']}-{i:02d}"
            med, p25, p75, n = report(ln)
            dflt, lo, hi, basis, flags = benchmark(ln, med, p25, p75, n)
            used = sum((usage[c][13] or 0) for c in ln["codes"] if c in usage)
            cust = max([usage[c][11] or 0 for c in ln["codes"] if c in usage] or [0])
            if ln["new"]:
                flags.append("NEW - no UNIS billing code yet")
            elif ln["tier"] == "Core" and used == 0 and ln["b2b"] is None and ln["d2c"] is None:
                flags.append("Core line with no Sep-26 invoice usage")
            rows.append(dict(lid=lid, section=title, sid=s["id"], service=s["name"], ln=ln, med=med, p25=p25, p75=p75, n=n,
                             default=dflt, lo=lo, hi=hi, basis=basis, flags=flags, used=used, cust=cust))
            for c in ln["codes"]:
                code_primary.setdefault(c, (lid, s["id"], s["name"], ln["name"], ln["tier"]))
        for c in s["extra"]:
            code_primary.setdefault(c, (f"{s['id']} (service)", s["id"], s["name"], "(service-level - no separate price line)", "Advanced"))
for grp, codes in spec.INTERNAL.items():
    for c in codes:
        code_primary.setdefault(c, ("INTERNAL", "INTERNAL", grp, "Billing adjustment - not on customer quote", "Internal"))

unmapped = sorted(set(master) - set(code_primary))
unknown = sorted(set(code_primary) - set(master))
assert not unmapped, f"unmapped codes: {unmapped}"
assert not unknown, f"unknown codes: {unknown}"

# ------------------------------------------------------------------ styles
F = "Calibri"
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
NEW = PatternFill("solid", fgColor="FFF2CC")
ADV = PatternFill("solid", fgColor="F2F2F2")
BAD = PatternFill("solid", fgColor="FCE4D6")
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row, i, c)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def put(ws, r, vals, fill=None, bold_cols=(), fmt=None, size=9):
    for col, v in enumerate(vals, 1):
        cell = ws.cell(r, col, v)
        cell.alignment = WRAP
        cell.border = BORDER
        cell.font = Font(name=F, size=size, bold=col in bold_cols)
        if fill:
            cell.fill = fill
        if fmt and col in fmt and isinstance(v, (int, float)):
            cell.number_format = fmt[col]

def money_fmt(ln):
    return "0%" if ln["pct"] else '"$"#,##0.00'

def sec_row(ws, r, title, ncols):
    ws.cell(r, 1, title).font = Font(name=F, bold=True, size=10)
    for c in range(1, ncols + 1):
        ws.cell(r, c).fill = SEC
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=min(ncols, 4))

def fmt_price(ln, v):
    if v is None:
        return "quote"
    return f"{v:.0%}" if ln["pct"] else f"${v:,.2f}"

wb = openpyxl.Workbook()
n_core = sum(1 for x in rows if x["ln"]["tier"] == "Core")
n_new = sum(1 for x in rows if x["ln"]["new"])
n_svc = sum(len(s) for _, s in spec.SECTIONS)

# ================================================================== README
ws = wb.active
ws.title = "README"
ws.column_dimensions["A"].width = 120
text = [
    ("UNIS Fulfillment - Rate Card Catalog v2 (industry-standard, webform-ready)", 14, True),
    ("Supersedes v1. Built from: Billing Items Usage Sep 2026 (318 usage codes + 16 master-only, 40,460 customer price-list rows), the Active-Items Scope file, the Northampton B2B and D2C 2025 templates, and general US 3PL industry practice.", 10, False),
    ("", 10, False),
    ("DECISIONS APPLIED (from your answers)", 11, True),
    ("1. Drayage excluded. The only yard-related items kept are warehouse-side charges (yard storage, yard hostler, outside-carrier container handling).", 10, False),
    (f"2. Granularity set by industry convention: {len(spec.SECTIONS)} sections > {n_svc} sub-items (services) > {len(rows)} price lines ({n_core} Core, {len(rows) - n_core} Advanced). Lines that describe the same work are ONE line; the price is a user-entered field.", 10, False),
    ("3. Lines whose codes were never invoiced, or that are niche, are tier = Advanced (hidden behind 'advanced options' on the webform).", 10, False),
    ("4. Every line has a benchmark default price and a range. Default = UNIS price-list median when it has >= 20 data points and sits inside the industry range; else the current UNIS template rate; else the industry midpoint. The basis is shown per line.", 10, False),
    ("5. One quote covers both channels. Each sub-item holds B2B, D2C and 'Both' lines side by side (column 'Channel'); the webform filters/labels by channel inside each section instead of producing two quotes.", 10, False),
    ("", 10, False),
    ("IMPORTANT - about the industry ranges", 11, True),
    ("The 'Industry range' columns are my estimates from general knowledge of US 3PL pricing (2025-26), NOT a cited market survey. Use them as a sanity band. Where UNIS's own data sits outside the band it is flagged. Please have commercial / ops review before customer use.", 10, False),
    ("", 10, False),
    ("TABS", 11, True),
    ("Scope & How-To - one row per sub-item in the reference layout: what we do / billing trigger & unit / rate-card line items (with channel tag and default price) / scope notes.", 10, False),
    ("Rate Card (Webform) - one row per price line = the data model for the Quotation webform (also exported as rate-card-catalog-v2.json).", 10, False),
    ("Review - Active Items File - row-by-row findings on UNIS_UF_Scope_How-To_Active_Items_Only.xlsx and what changed.", 10, False),
    ("Charge Code Mapping - all 334 system codes -> the v2 price line that bills them (Core / Advanced / Internal).", 10, False),
    ("Colour key: yellow = NEW industry-standard line with no UNIS billing code yet; grey = Advanced; orange = issue.", 10, False),
]
for i, (t, sz, b) in enumerate(text, 1):
    c = ws.cell(i, 1, t)
    c.font = Font(name=F, size=sz, bold=b)
    c.alignment = WRAP

# ================================================================== SCOPE & HOW-TO
ws = wb.create_sheet("Scope & How-To", 0)
cols = ["Activity - How We Do It", "Billing Trigger / Unit", "Rate Card Line Item(s)  [channel] default price / unit", "Scope Notes", "ID", "Channels", "Lines (Core / Adv)"]
ws.cell(1, 1, "UNIS Fulfillment - Scope of Services & Billing Methodology (v2)").font = Font(name=F, bold=True, size=14)
ws.merge_cells("A1:D1")
c = ws.cell(2, 1, "Plain-language reference of what each service covers and how it is billed. Each service lists its rate-card lines with the channel it applies to (B2B, D2C or Both) and the benchmark default price; lines marked (adv) are advanced options. Detailed ranges and data evidence: 'Rate Card (Webform)' tab.")
c.font = Font(name=F, size=9, color="595959"); c.alignment = WRAP
ws.merge_cells("A2:D2"); ws.row_dimensions[2].height = 40
header(ws, 4, cols, [50, 30, 62, 46, 8, 11, 10])
ws.freeze_panes = "A5"
r = 5
by_sid = collections.defaultdict(list)
for x in rows:
    by_sid[x["sid"]].append(x)
for title, services in spec.SECTIONS:
    sec_row(ws, r, title, len(cols)); r += 1
    for s in services:
        ls = by_sid[s["id"]]
        items = "\n".join(f"- [{x['ln']['ch']}] {x['ln']['name']}: {fmt_price(x['ln'], x['default'])} / {x['ln']['unit']}"
                          + (" (adv)" if x["ln"]["tier"] == "Advanced" else "") + (" *NEW*" if x["ln"]["new"] else "") for x in ls)
        chans = sorted({x["ln"]["ch"] for x in ls})
        core = sum(1 for x in ls if x["ln"]["tier"] == "Core")
        put(ws, r, [s["what"], s["trigger"], f"{s['name']}\n{items}", s["notes"] or "", s["id"], " / ".join(chans), f"{core} / {len(ls) - core}"], bold_cols=(5,))
        r += 1

# ================================================================== RATE CARD (WEBFORM)
ws = wb.create_sheet("Rate Card (Webform)", 1)
cols = ["Line ID", "Section", "Service ID", "Service", "Line Item", "What's included / when it applies", "Unit", "Channel", "Tier",
        "Default Price (benchmark)", "Range Low", "Range High", "Default Basis",
        "UNIS Median", "UNIS P25", "UNIS P75", "UNIS # Rates", "UNIS Template B2B", "UNIS Template D2C",
        "Sep-26 Invoice Lines", "Flags", "Charge Codes"]
header(ws, 1, cols, [10, 22, 8, 26, 34, 40, 13, 8, 9, 11, 9, 9, 26, 9, 9, 9, 8, 9, 9, 9, 28, 34])
ws.freeze_panes = "F2"
for i, x in enumerate(rows, 2):
    ln = x["ln"]
    mf = money_fmt(ln)
    fill = NEW if ln["new"] else ADV if ln["tier"] == "Advanced" else None
    put(ws, i, [x["lid"], x["section"].split(". ", 1)[1].title(), x["sid"], x["service"], ln["name"], ln["desc"], ln["unit"], ln["ch"], ln["tier"],
                x["default"], x["lo"], x["hi"], x["basis"], x["med"], x["p25"], x["p75"], x["n"], ln["b2b"], ln["d2c"], x["used"],
                "; ".join(x["flags"]), ", ".join(ln["codes"])],
        fill=fill, bold_cols=(10,), fmt={k: mf for k in (10, 11, 12, 14, 15, 16, 18, 19)})
    if x["flags"] and any("outside" in f for f in x["flags"]):
        ws.cell(i, 21).fill = BAD
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

# ================================================================== REVIEW OF ACTIVE-ITEMS FILE
ws = wb.create_sheet("Review - Active Items File", 2)
af = openpyxl.load_workbook(f"{IN}/UNIS_UF_Scope_How-To_Active_Items_Only.xlsx")["Scope & How-To"]
merged_rows = {rng.min_row for rng in af.merged_cells.ranges}
MERGED_GUESS = {16: "HANDLING-0178", 17: "HANDLING-0174", 23: "HANDLING-0109 / 0111 / 0125", 24: "HANDLING-0109 / 0111 / 0125",
                28: "HANDLING-0109 / 0111 / 0125", 29: "HANDLING-0120", 36: "HANDLING-0081", 37: "HANDLING-0052",
                41: "duplicate of HANDLING-0043 / 0044?", 42: "HANDLING-0029"}
SPECIFIC = {
    "ACCESSORIAL-0038": "Code is the generic 'Material' charge (not in the usage sheet). Two of its three rows reuse the Stretch Wrap and Palletizing & Wrapping text - wrong code for those services.",
    "ACCESSORIAL-0001": "Code is cancel BEFORE picking, but the text says 'after chargeable work has started'; scope note is copied from a picking row ('Use for Order-level picking').",
    "ACCESSORIAL-0002": "Scope note copied from a picking row ('Use for Order-level picking') - should state the cancellation cut-off and return-to-stock labor.",
    "ACCESSORIAL-0007": "Same service as ACCESSORIAL-0002 (cancel after pick) - merge.",
    "ACCESSORIAL-0018": "This is a non-compliance fee (client did not send the ASN / packing list in time), not 'prepare documentation'.",
    "HANDLING-0095": "Outbound missed ORDER appointment is placed under INBOUND and described as 'Receive inbound product'.",
    "HANDLING-0041": "Inbound receipt processing placed under OUTBOUND ('outbound warehouse activity').",
    "HANDLING-0042": "Inbound receipt processing placed under OUTBOUND ('outbound warehouse activity').",
    "HANDLING-0186": "Assembly put-away note copied from Put Away ('Applies after inbound receiving'); it is put-away of finished kits.",
    "HANDLING-0188": "Receiving charge (CHEP truckload offload) filed under Accessorials.",
    "HANDLING-0252": "Receiving charge (truckload offload) filed under Accessorials.",
    "STORAGE INCOME-0027": "Yard / container storage filed under Accessorials.",
    "STORAGE INCOME-0028": "Space rental is storage, filed under Accessorials.",
    "SYSTEM & MANAGEMENT FEE-0005": "IT / system fee filed under Accessorials - give IT its own section.",
    "SYSTEM & MANAGEMENT FEE-0003": "IT / system fee filed under Accessorials; duplicate of SYSTEM & MANAGEMENT FEE-0005.",
    "HANDLING-0194": "IT project fee filed under Accessorials - give IT its own section.",
    "HANDLING-0217": "IT support filed under Accessorials - give IT its own section.",
    "ACCESSORIAL-0004": "Claim = billing adjustment, not a service - remove from customer scope.",
    "HANDLING-0198": "Claim = billing adjustment, not a service - remove from customer scope.",
    "ACCESSORIAL-0003": "Generic catch-all code with no defined scope - should not be on a customer rate card.",
    "HANDLING-0254": "Generic catch-all code with no defined scope - should not be on a customer rate card.",
    "HANDLING-0195": "Generic 'Handling In' - duplicate of receiving lines; no defined scope.",
    "HANDLING-0197": "Generic 'Handling Out' - duplicate of picking/loading lines; no defined scope.",
    "HANDLING-0191": "Generic 'Customer Service' - fold into account management fee.",
    "HANDLING-0093": "Same code appears in both the B2B and DTC sections. OK as a channel split - v2 keeps one sub-item with a B2B line and a D2C line.",
    "HANDLING-0072": "Good anti-double-billing note ('do not also bill the standard case pick') - kept in v2.",
    "STORAGE INCOME-0004": "Good note on the combined Initial + Recurring total minimum - adopted in v2 (ST-01 / ST-08).",
    "HANDLING-0025": "Transload is not receiving-to-stock - moved to 'Transload & Cross-Dock'.",
    "HANDLING-0023": "Transload is not receiving-to-stock - moved to 'Transload & Cross-Dock'.",
    "HANDLING-0008": "Transload is not receiving-to-stock - moved to 'Transload & Cross-Dock'.",
    "HANDLING-0010": "Transload is not receiving-to-stock - moved to 'Transload & Cross-Dock'.",
    "HANDLING-0011": "Transload is not receiving-to-stock - moved to 'Transload & Cross-Dock'.",
}
review = []
seen_codes = collections.Counter()
section_name = None
for row in range(5, af.max_row + 1):
    a, b, c3, d = (af.cell(row, k).value for k in range(1, 5))
    if a and not any((b, c3, d)) and row not in MERGED_GUESS:
        section_name = a
        continue
    if not a:
        continue
    issues, codes = [], []
    if row in MERGED_GUESS:
        issues.append("Row is merged across A:D - billing trigger, code and scope notes are lost.")
        codes = [f"probably {MERGED_GUESS[row]}"]
    else:
        m = re.match(r"([A-Z &]+-\d{4})(?: / (\d{4}))?", str(c3))
        if m:
            codes = [m.group(1)] + ([m.group(1).rsplit("-", 1)[0] + "-" + m.group(2)] if m.group(2) else [])
    for code in codes:
        if code.startswith("probably"):
            continue
        seen_codes[code] += 1
        if code not in usage:
            issues.append("Code is not in the usage sheet (no invoice usage) - contradicts the file's 'active items only' rule.")
        if code in SPECIFIC:
            issues.append(SPECIFIC[code])
    if isinstance(a, str) and a.startswith("Perform the separately contracted"):
        issues.append("Boilerplate activity text - does not say what work is done.")
    elif isinstance(a, str) and re.search(r"required (receiving|outbound warehouse) activity for|complete the applicable receiving handling", a):
        issues.append("Activity text repeats the system name instead of describing the work.")
    if isinstance(d, str) and d.startswith("Use when this service is separately included"):
        issues.append("Generic scope note - does not state what is included / excluded.")
    probable = [f"HANDLING-{n}" for x in codes if x.startswith("probably") for n in re.findall(r"\b0\d{3}\b", x)]
    v2 = "; ".join(sorted({code_primary[c][0] for c in codes + probable if c in code_primary}))
    review.append((row, section_name, codes, a, b, issues, v2))

# summary
missing_active = sorted([c for c, u in usage.items() if (u[13] or 0) > 0 and c not in seen_codes and
                         not any(c in g for g in MERGED_GUESS.values())], key=lambda c: -(usage[c][13] or 0))
not_active = [c for c in seen_codes if c not in usage]
dups = [c for c, k in seen_codes.items() if k > 1]
summary = [
    ("Summary", f"{len(review)} item rows reviewed; {sum(1 for x in review if x[5])} have at least one finding."),
    ("1. Lost rows", f"{len(MERGED_GUESS)} rows are merged across A:D (likely a formatting slip) so their unit, code and notes are gone - these are the rows whose codes appear 'missing'."),
    ("2. Filter not applied", f"{len(not_active)} codes in the file are not in the usage sheet at all (0 invoice usage): {', '.join(not_active)}."),
    ("3. Active codes missing", f"{len(missing_active)} codes WITH invoice usage are absent (excl. the merged rows): " + ", ".join(f"{c} ({usage[c][13]} lines)" for c in missing_active) + "."),
    ("4. Duplicates", f"{', '.join(dups)} appear more than once (ACCESSORIAL-0038 x3 with wrong descriptions)."),
    ("5. Mis-filed / wrong text", "Missed-order appointment under Inbound; receipt processing under Outbound; transload under Receiving; receiving, storage, yard and IT items under Accessorials; cancel-before-pick described as after-pick; No-ASN described as a document service."),
    ("6. Granularity", f"One row per system code ({len(review)} rows): the same service repeats per UOM / pricing method (11 offload rows, 14 pick rows, 7 order-processing rows). A customer cannot compare these. v2 merges them into sub-items with price lines."),
    ("7. Plain language", "About 45 rows use boilerplate ('Perform the separately contracted X service' / 'Use when this service is separately included...'). v2 rewrites every line with what is included and when it applies."),
    ("8. Industry gaps", f"Missing standard 3PL items (added in v2, yellow): account setup, EDI per partner, e-commerce connector, pallets A/B, corner boards, branded packaging, inserts, fragile wrap, gift wrap, poly-bag prep, weight surcharge 10-30 lb, lot/expiry capture, long-term storage surcharge, peak surcharge, account minimum ({n_new} lines)."),
    ("9. Keep", "Good ideas in the file adopted in v2: anti-double-billing notes (weight-tier vs standard pick; palletize & wrap vs stretch wrap) and the combined Initial + Recurring storage minimum."),
]
header(ws, 1, ["Topic", "Finding"], [22, 150])
r = 2
for t, f in summary:
    put(ws, r, [t, f], bold_cols=(1,), size=10); r += 1
r += 1
cols = ["File Row", "Section in file", "Code(s)", "Activity text (file)", "Trigger / Unit (file)", "Findings", "Maps to v2 line"]
for i, cname in enumerate(cols, 1):
    cell = ws.cell(r, i, cname)
    cell.font = Font(name=F, bold=True, color="FFFFFF", size=10); cell.fill = HDR; cell.border = BORDER
for i, w in enumerate([8, 22, 24, 50, 26, 60, 22], 1):
    ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, w)
r += 1
for row, secn, codes, a, b, issues, v2 in review:
    put(ws, r, [row, secn, ", ".join(codes), a, b, "\n".join(f"- {x}" for x in issues) or "OK", v2], fill=BAD if issues else None)
    r += 1
ws.freeze_panes = ws.cell(len(summary) + 4, 1)

# ================================================================== CODE MAPPING
ws = wb.create_sheet("Charge Code Mapping")
cols = ["Charge Code", "System Item Name", "System Description", "UOM", "Category", "v2 Line ID", "v2 Service", "v2 Line Item", "Tier",
        "Customers", "Sep-26 Invoice Lines", "Status"]
header(ws, 1, cols, [26, 36, 40, 12, 16, 14, 30, 38, 9, 9, 10, 24])
ws.freeze_panes = "B2"
def sortkey(c):
    lid = code_primary[c][0]
    return (lid == "INTERNAL", lid, -((usage.get(c) or [0] * 14)[13] or 0))
for i, c in enumerate(sorted(master, key=sortkey), 2):
    lid, sid, sname, lname, tier = code_primary[c]
    if c in usage:
        u = usage[c]
        vals = [c, u[3], u[4], u[5], u[2], lid, sname, lname, tier, u[11], u[13], "In use" if (u[13] or 0) > 0 else "Configured - not invoiced"]
    else:
        m = master[c]
        vals = [c, m[2], "", m[3], m[6], lid, sname, lname, tier, None, None, "Master list only"]
    put(ws, i, vals, fill=ADV if tier == "Advanced" else BAD if tier == "Internal" else None)
ws.auto_filter.ref = f"A1:L{len(master) + 1}"

for w in wb.worksheets:
    w.sheet_view.zoomScale = 110
wb.move_sheet("README", offset=-wb.index(wb["README"]))
wb.active = 1
wb.save(OUT)

# ------------------------------------------------------------------ JSON for the webform (Task 2)
catalog = {"version": "v2", "generated_from": "Billing Items Usage Sep 2026", "sections": []}
for title, services in spec.SECTIONS:
    sec = {"title": title.split(". ", 1)[1], "services": []}
    for s in services:
        sec["services"].append({"id": s["id"], "name": s["name"], "what": s["what"], "trigger": s["trigger"], "notes": s["notes"],
            "lines": [{"id": x["lid"], "name": x["ln"]["name"], "desc": x["ln"]["desc"], "unit": x["ln"]["unit"], "channel": x["ln"]["ch"],
                       "tier": x["ln"]["tier"], "isPercent": x["ln"]["pct"], "default": x["default"], "rangeLow": x["lo"], "rangeHigh": x["hi"],
                       "basis": x["basis"], "isNew": x["ln"]["new"], "chargeCodes": x["ln"]["codes"]} for x in by_sid[s["id"]]]})
    catalog["sections"].append(sec)
json.dump(catalog, open(JSON_OUT, "w"), indent=1)

print(f"saved {OUT}\nsections {len(spec.SECTIONS)} services {n_svc} lines {len(rows)} core {n_core} new {n_new}")
for x in rows:
    if x["flags"] and not x["ln"]["new"]:
        print("  FLAG", x["lid"], x["ln"]["name"], "|", "; ".join(x["flags"]))
