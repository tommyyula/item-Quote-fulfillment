"""Single loader for the billing-item source workbook, read by column NAME (layouts differ between exports).

Default source: workingfolder-input/BillingItem_Usage1001 -Final.xlsx (override with env RATE_SOURCE).
Exposes:
  usage[code]   -> dict(name, desc, uom, category, cond, customers, hl_customers, freq, trigger)
  master[code]  -> dict(name, uom, category, trigger)
  price_rows    -> list of dicts (customer, code, name, desc, rate_type, rate)   - all customers' price lists
  common        -> [dict(code, name, uom, tag, hl_customers)] from 'Final Common Billing Items' (may be empty)
"""
import os
import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
SOURCE = os.environ.get("RATE_SOURCE", f"{ROOT}/workingfolder-input/BillingItem_Usage1001 -Final.xlsx")
_wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)


def _sheet(name):
    for s in _wb.sheetnames:
        if s.strip().lower() == name.lower():
            return _wb[s]
    return None


def _records(name):
    ws = _sheet(name)
    if ws is None:
        return []
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        if not any(v is not None for v in r):
            continue
        rec = {}
        for h, v in zip(head, r):
            if h and h not in rec:  # first column of a repeated header (e.g. two 'Customer Name' columns) wins
                rec[h] = v
        out.append(rec)
    return out


def _num(x):
    try:
        return float(str(x).replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None


usage = {}
for r in _records("All Billing Items Usage"):
    code = r.get("Charge Code")
    if not code:
        continue
    usage[code] = dict(name=r.get("Item Name") or "", desc=r.get("Item Description") or "", uom=r.get("UOM") or "",
                       category=r.get("Category") or "", cond=r.get("Condition"),
                       customers=int(_num(r.get("Customer Count by Invoice")) or 0),
                       hl_customers=int(_num(r.get("High Level Customer Count")) or 0),
                       freq=int(_num(r.get("Usage Frequency in the invoice")) or 0), trigger=r.get("TriggerPoint"))

master = {}
for r in _records("All Billing Items"):
    code = r.get("ChargeCode")
    if code:
        master[code] = dict(name=r.get("ItemName") or "", uom=r.get("UOM") or "", category=r.get("Category") or "", trigger=r.get("TriggerPoint"))

price_rows = [dict(customer=r.get("CustomerCode"), customer_name=r.get("CustomerName"), code=r.get("Charge Code"), name=r.get("Item Name"),
                   desc=str(r.get("Description") or "").strip(), rate_type=r.get("Rate Type") or "", rate=_num(r.get("Rate")))
              for r in _records("All the customers price list") if r.get("Charge Code")]

common = [dict(code=r["Charge Code"], name=r.get("Item Name") or "", uom=r.get("UOM") or "", tag=r.get("Tag") or "",
               hl_customers=int(_num(r.get("High Level Customer Count")) or 0))
          for r in _records("Final Common Billing Items") if r.get("Charge Code")]
